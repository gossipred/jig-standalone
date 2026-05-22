from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "JJ-Jig-Management-System-Java-Project-Report.docx"


def set_run(run, size=11, bold=False, color="17202A"):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Jig & Toolings Management System")
    set_run(r, 24, True, "0F172A")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("治具及模具管理系統")
    set_run(r, 18, True, "334155")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Java 程式設計結訓專題報告")
    set_run(r, 14, False, "475569")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("作者：JJ  |  技術：Java 17, Spring Boot, MySQL, Spring Security")
    set_run(r, 11, False, "64748B")
    doc.add_paragraph()


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.color.rgb = RGBColor.from_string("0F172A")
    return p


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run(r)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run(run, 10, True, "FFFFFF")
        shading = hdr[i]._tc.get_or_add_tcPr()
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "101820")
        shading.append(shd)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for run in p.runs:
                    set_run(run, 9, False, "17202A")
    doc.add_paragraph()


def add_figure(doc, image_name, caption):
    image_path = ROOT / "ppt-preview" / image_name
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(image_path), width=Inches(6.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cap.add_run(caption)
    set_run(rr, 9, False, "64748B")


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(11)

    add_title(doc)

    h(doc, "1. 專題動機")
    para(doc, "本專題以工廠治具與模具管理為情境，將現場常見的 Excel、紙本、個人資料夾與附件分散問題，轉換為一套 Java Spring Boot Web 系統。")
    bullets(doc, [
        "治具/模具資料不易快速查詢。",
        "圖片、圖紙、PDF、CAD 檔案常分散存放。",
        "狀態與到期日變更缺乏完整追蹤。",
        "不同角色權限不同，需要避免誤改與誤刪。",
    ])

    h(doc, "2. 專題目標")
    para(doc, "本系統目標是完成一套可登入、可管理主資料、可上傳附件、可追蹤異動、可依角色控管權限的 Java Web 專題。")
    add_figure(doc, "slide-03.png", "Figure 1. Project goals: master data, files, roles, logs, and handoff.")

    h(doc, "3. 系統架構")
    para(doc, "系統採用 Spring Boot MVC 架構。Controller 接收請求，Service 處理商業邏輯，Repository 透過 JPA 存取 MySQL。畫面使用 Thymeleaf 與 Bootstrap，登入權限使用 Spring Security。")
    add_figure(doc, "slide-04.png", "Figure 2. Spring Boot MVC architecture.")

    h(doc, "4. 技術堆疊")
    add_table(doc, ["Layer", "Technology", "Purpose"], [
        ["Language", "Java 17", "主要開發語言"],
        ["Framework", "Spring Boot", "建立 Web 系統與整合功能"],
        ["Build", "Maven", "管理 dependency 與測試流程"],
        ["Database", "MySQL 8.4", "正式資料庫"],
        ["ORM", "Spring Data JPA / Hibernate", "Entity 與資料庫存取"],
        ["Security", "Spring Security", "登入與角色權限"],
        ["View", "Thymeleaf + Bootstrap", "網頁畫面與互動介面"],
    ])

    h(doc, "5. 資料庫設計")
    para(doc, "資料庫設計重點是將目前狀態與歷史紀錄分開保存。主資料表保存目前資料，log table 保存異動歷程。")
    add_figure(doc, "slide-06.png", "Figure 3. Database tables and traceability relationships.")

    h(doc, "6. 角色權限")
    add_table(doc, ["Role", "Best For", "Allowed", "Restriction"], [
        ["ADMIN", "系統管理者", "帳號管理、核准報廢、修改所有資料", "第一版不限制"],
        ["SUPERVISOR", "主管/線長", "跨使用者修改與刪除、看紀錄、匯出", "不可管理帳號"],
        ["ENGINEER", "工程/維修人員", "新增與修改自己建立的資料", "不可刪除或報廢"],
        ["OPERATOR", "產線查詢人員", "查詢、下載、查看紀錄、匯出", "唯讀"],
    ])
    add_figure(doc, "slide-07.png", "Figure 4. Role permission matrix.")

    h(doc, "7. 主要功能")
    bullets(doc, [
        "Jig / Tooling List：以 Classification 作為第一欄，支援多欄位關鍵字查詢。",
        "Add / Edit Jig：建立與維護治具/模具主資料。",
        "Jig No. 半自動給號：預設 AM-ME，可調整前綴，儲存時檢查重號。",
        "檔案上傳下載：每筆資料最多 5 個附件。",
        "Status / Due Date：支援狀態與到期日更新。",
        "Change Logs：記錄使用者、時間、Old / New 與備註。",
        "User Management：ADMIN 管理帳號，使用者可自行改密碼。",
    ])
    add_figure(doc, "slide-09.png", "Figure 5. Main list UI concept.")
    add_figure(doc, "slide-11.png", "Figure 6. Change log old/new tracking.")

    h(doc, "8. 開發歷程與需求修正")
    para(doc, "專題開發過程中依照實際操作逐步修正需求。例如從 Jig 擴充為 Jig & Toolings，URL 欄位改成檔案上傳，新增 SUPERVISOR，並補強異動紀錄 old/new。")
    add_figure(doc, "slide-10.png", "Figure 7. Iterative development changes.")

    h(doc, "9. 測試與驗證")
    para(doc, "本專題使用 Spring Boot Test 與 H2 測試資料庫進行自動化測試。最新測試結果為 47 tests passed, 0 failures, 0 errors。")
    add_figure(doc, "slide-12.png", "Figure 8. Automated testing result.")

    h(doc, "10. 部署與客戶移交規劃")
    para(doc, "客戶移交時，建議使用一台內部網路電腦作為 server，安裝 Java、MySQL 與系統 JAR。其他使用者只要用瀏覽器連線到 server IP 即可使用。")
    add_figure(doc, "slide-13.png", "Figure 9. Internal-network deployment concept.")

    h(doc, "11. 結論")
    para(doc, "Jig & Toolings Management System 將工程管理需求轉化為完整 Java Web 系統，涵蓋登入權限、資料庫設計、主資料管理、附件、異動紀錄、測試與移交規劃。這份專題不只是 Java 語法練習，也呈現了實務系統開發的完整流程。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

