from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "治具及模具管理系統_中文移交安裝說明.docx"


def set_east_asia_font(run, font_name="Microsoft JhengHei"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9.5)
    set_east_asia_font(r)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    set_east_asia_font(r)
    return p


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(11)
    set_east_asia_font(r)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(45, 55, 72)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        r.font.size = Pt(10.5)
        set_east_asia_font(r)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
    styles[name].font.name = "Microsoft JhengHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
styles["Normal"].font.size = Pt(11)
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.color.rgb = RGBColor(46, 116, 181)
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.color.rgb = RGBColor(46, 116, 181)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("治具及模具管理系統")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(11, 37, 69)
set_east_asia_font(r)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Jig & Toolings Management System\n中文移交安裝說明")
r.font.size = Pt(14)
set_east_asia_font(r)

add_body(doc, "文件版本：2026-05-22")
add_body(doc, "適用環境：Windows 10 / Windows 11，作為公司內部伺服器使用")

add_heading(doc, "1. 移交方式總覽")
add_body(doc, "這套系統建議安裝在一台固定開機的 Windows 電腦上，當作公司內部伺服器。其他使用者不用安裝程式，只要用瀏覽器連到主機 IP 即可使用。")
add_code(doc, "Windows 主機 → Java 17 + MySQL → 啟動系統 → 使用者用 http://主機IP:8080 連線")

add_heading(doc, "2. 客戶需要準備什麼")
add_table(
    doc,
    ["項目", "建議", "說明"],
    [
        ["作業系統", "Windows 10 / 11", "一般辦公室電腦即可"],
        ["Java", "Java 17", "系統主程式需要 Java 17 執行"],
        ["資料庫", "MySQL Server 8.x", "儲存帳號、治具、模具、異動紀錄"],
        ["瀏覽器", "Chrome / Edge", "使用者操作系統"],
        ["網路", "公司內部 LAN", "其他同事透過 IP 連線"],
        ["備份空間", "外接硬碟或公司備份磁碟", "備份資料庫與上傳檔案"],
    ],
    [1.2, 1.6, 3.7],
)

add_heading(doc, "3. 移交包內容")
add_table(
    doc,
    ["位置", "用途", "客戶重點"],
    [
        ["app/jig-toolings-management.jar", "系統主程式", "不用打開，由 bat 啟動"],
        ["app/application.properties", "系統設定", "設定 MySQL 密碼、port、上傳路徑"],
        ["app/start-system.bat", "啟動系統", "每次要開系統就點這個"],
        ["app/stop-system.bat", "停止系統", "需要關閉系統時使用"],
        ["database/install-database.bat", "安裝資料庫", "第一次安裝時執行"],
        ["scripts/backup-now.bat", "手動備份", "備份資料庫與上傳檔案"],
        ["uploads/jigs", "附件存放", "圖片、圖紙、PDF、CAD 檔案"],
    ],
    [2.3, 1.35, 2.85],
)

add_heading(doc, "4. 第一次安裝流程")
steps = [
    "複製整個移交包到 C:\\JJ-Jig-Toolings\\，不要只複製其中一部分。",
    "安裝 Java 17，並用 java -version 確認版本包含 17。",
    "安裝 MySQL Server 8.x，記住 root 密碼。",
    "用記事本打開 app\\application.properties，把 MySQL 密碼寫在 spring.datasource.password= 後面，按 Ctrl + S 儲存後關閉記事本。",
    "執行 database\\install-database.bat，建立資料庫與預設帳號。",
    "執行 app\\start-system.bat，啟動系統並保持黑色視窗開啟。",
    "主機瀏覽器開 http://localhost:8080。",
    "其他同事用 http://主機IP:8080 連線，例如 http://192.168.1.50:8080。",
]
for i, step in enumerate(steps, 1):
    add_body(doc, f"{i}. {step}")

add_heading(doc, "4-1. MySQL 密碼要寫在哪裡")
add_body(doc, "MySQL 安裝時設定的 root 密碼，要寫在 app\\application.properties 裡。這不是系統登入密碼，而是系統連接資料庫用的密碼。")
add_body(doc, "操作方式：進入移交包資料夾 → 打開 app 資料夾 → 在 application.properties 上按右鍵 → 開啟方式 → 記事本。")
add_body(doc, "找到以下兩行：")
add_code(doc, "spring.datasource.username=root\nspring.datasource.password=")
add_body(doc, "密碼要寫在 spring.datasource.password= 後面，中間不要加空格。")
add_body(doc, "如果 MySQL root 密碼是 MyPassword123，請改成：")
add_code(doc, "spring.datasource.username=root\nspring.datasource.password=MyPassword123")
add_body(doc, "如果客戶 IT 建立的是其他 MySQL 帳號，例如 jigadmin，請同時修改 username 與 password：")
add_code(doc, "spring.datasource.username=jigadmin\nspring.datasource.password=客戶設定的MySQL密碼")
add_body(doc, "修改完成後，按 Ctrl + S 儲存，然後按記事本右上角 X 關閉。如果跳出是否儲存變更，請選儲存。")
add_body(doc, "注意：不要把密碼寫到別的地方，也不要刪除等號前面的文字。正確位置只有 spring.datasource.password=你的MySQL密碼。", bold=True)

add_heading(doc, "5. 預設帳號與權限")
add_body(doc, "預設密碼全部為 123456。第一次登入後，請立刻修改密碼。", bold=True)
add_table(
    doc,
    ["帳號", "權限", "用途"],
    [
        ["admin", "ADMIN", "系統管理者，可管理帳號與所有資料"],
        ["supervisor", "SUPERVISOR", "主管，可修改與刪除所有治具/模具資料"],
        ["engineer", "ENGINEER", "工程人員，可新增與修改自己建立的資料"],
        ["operator", "OPERATOR", "操作人員，只能查詢、下載、匯出"],
    ],
    [1.4, 1.6, 3.5],
)

add_heading(doc, "6. 日常使用與備份")
add_body(doc, "一般使用流程：登入 → 搜尋 Jig / Tooling → 查看詳細資料 → 依權限更新狀態或到期日 → 上傳或下載圖紙檔案 → 查看異動紀錄 → 匯出 CSV。")
add_body(doc, "備份請執行 scripts\\backup-now.bat。備份內容包含 MySQL 資料庫，以及 uploads\\jigs 裡的圖片、圖紙、PDF、CAD 檔案。")
add_bullets(doc, ["每天多人使用：建議每天下班前備份。", "偶爾使用：建議每週備份。", "正式量產環境：建議交由客戶 IT 做自動備份。"])

add_heading(doc, "7. 內網與防火牆")
add_body(doc, "如果主機可以開 http://localhost:8080，但其他電腦不能開 http://主機IP:8080，通常是 Windows 防火牆擋住 8080 port。")
add_body(doc, "請用系統管理員身分執行 scripts\\open-firewall-8080.bat。")

add_heading(doc, "8. 常見問題")
add_table(
    doc,
    ["問題", "可能原因", "處理方式"],
    [
        ["打不開 localhost:8080", "系統沒有啟動", "執行 app\\start-system.bat"],
        ["資料庫錯誤", "MySQL 密碼錯誤", "修改 app\\application.properties"],
        ["其他電腦連不上", "防火牆擋住 8080", "執行 open-firewall-8080.bat"],
        ["java 不是內部或外部命令", "Java 未安裝或 PATH 未設定", "重新安裝 Java 17"],
        ["上傳檔案找不到", "uploads/jigs 被移動或刪除", "還原備份資料夾"],
        ["忘記密碼", "使用者無法自行找回", "由 admin 到 User Management 重設"],
    ],
    [1.85, 1.9, 2.75],
)

add_heading(doc, "9. 交接檢查表")
check_items = [
    "客戶已指定 Windows 主機",
    "已確認主機 IP",
    "Java 17 已安裝",
    "MySQL Server 8.x 已安裝",
    "application.properties 已設定 MySQL 密碼",
    "install-database.bat 已執行成功",
    "start-system.bat 可啟動系統",
    "主機可開啟 http://localhost:8080",
    "其他電腦可開啟 http://主機IP:8080",
    "admin 可登入並新增使用者",
    "可新增 Jig / Tooling",
    "可上傳與下載檔案",
    "可匯出 CSV",
    "異動紀錄有記錄時間、使用者、動作、Old/New",
    "客戶知道如何備份",
    "客戶已修改預設密碼",
]
for item in check_items:
    add_body(doc, f"☐ {item}")

add_heading(doc, "10. 建議交付說法")
add_body(doc, "這套系統安裝在一台 Windows 主機上，其他使用者不用安裝程式，只要用瀏覽器連到主機 IP 就能使用。資料存在 MySQL，圖片與圖紙存在 uploads 資料夾。每天只要備份資料庫和 uploads，就能保留完整資料。")
add_body(doc, "平常不用打開程式碼，也不用開 VS Code。日常只需要會點 start-system.bat、登入系統、做備份，就能使用。")

doc.save(OUT)
print(OUT)
